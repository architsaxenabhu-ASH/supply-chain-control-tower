"""Payables Engine (Phase 3D, P1).

Mirrors the receivables architecture for money we owe partners (suppliers, 3PLs,
freight forwarders, customs brokers, warehouse partners). Same status rules
(open / partially_paid / paid / overdue), payment history, and Excel upload.
Every write is audited.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.db.local_persistence import load_collection, record_audit_event, save_collection
from app.schemas.payables import PARTNER_TYPES, CreatePayableRequest, Payable
from app.schemas.receivables import PaymentEntry, ReceivableUploadSummary, RecordPaymentRequest


def _load() -> list[Payable]:
    return load_collection("payables", lambda payload: Payable(**payload))


def _save(records: list[Payable]) -> None:
    save_collection("payables", records, lambda record: record.payable_id)


def _next_id(existing: list[Payable]) -> str:
    numbers = [
        int(p.payable_id.split("-")[-1])
        for p in existing
        if p.payable_id.startswith("PAY-") and p.payable_id.split("-")[-1].isdigit()
    ]
    return f"PAY-{(max(numbers) + 1) if numbers else 1:04d}"


def _effective_status(record: Payable, today: date | None = None) -> str:
    today = today or date.today()
    if record.invoice_value - record.paid_value <= 0:
        return "paid"
    try:
        if date.fromisoformat(record.due_date) < today:
            return "overdue"
    except ValueError:
        pass
    return "partially_paid" if record.paid_value > 0 else "open"


def _enrich(record: Payable) -> Payable:
    record.outstanding_value = max(record.invoice_value - record.paid_value, 0)
    record.status = _effective_status(record)
    return record


def create_payable(request: CreatePayableRequest) -> Payable:
    partner_type = request.partner_type.strip().lower()
    if partner_type not in PARTNER_TYPES:
        raise ValueError(f"Partner type must be one of {PARTNER_TYPES}.")
    if request.invoice_value <= 0:
        raise ValueError("Invoice value must be greater than zero.")
    records = _load()
    record = Payable(
        payable_id=_next_id(records),
        partner_type=partner_type,
        partner_name=request.partner_name.strip(),
        country=request.country.strip(),
        invoice_number=request.invoice_number.strip(),
        invoice_date=request.invoice_date,
        due_date=request.due_date,
        payment_terms=request.payment_terms,
        invoice_value=request.invoice_value,
        paid_value=request.paid_value,
        created_by=request.actor,
        updated_at=datetime.now().isoformat(),
    )
    _save([record, *records])
    record_audit_event(
        action="payable_create",
        module_name="payables",
        entity_name="payable",
        entity_id=record.payable_id,
        actor=request.actor,
        new_value=record,
    )
    return _enrich(record)


def list_payables(partner_name: str | None = None, partner_type: str | None = None, status: str | None = None) -> list[Payable]:
    records = [_enrich(p) for p in _load()]
    if partner_name:
        records = [p for p in records if p.partner_name.lower() == partner_name.lower()]
    if partner_type:
        records = [p for p in records if p.partner_type == partner_type.lower()]
    if status:
        records = [p for p in records if p.status == status.lower()]
    return sorted(records, key=lambda p: p.invoice_date, reverse=True)


def get_payable(payable_id: str) -> Payable | None:
    record = next((p for p in _load() if p.payable_id == payable_id), None)
    return _enrich(record) if record else None


def record_payment(payable_id: str, request: RecordPaymentRequest) -> Payable:
    if request.amount <= 0:
        raise ValueError("Payment amount must be greater than zero.")
    records = _load()
    record = next((p for p in records if p.payable_id == payable_id), None)
    if not record:
        raise ValueError(f"Payable not found: {payable_id}")
    record.paid_value = min(record.invoice_value, record.paid_value + request.amount)
    record.payment_history.append(
        PaymentEntry(amount=request.amount, paid_date=request.paid_date or date.today().isoformat(), note=request.note, actor=request.actor)
    )
    record.updated_at = datetime.now().isoformat()
    _save(records)
    record_audit_event(
        action="payable_payment",
        module_name="payables",
        entity_name="payable",
        entity_id=payable_id,
        actor=request.actor,
        new_value=record,
    )
    return _enrich(record)


def _cell(row: dict, *names):
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def upload_payables(content: bytes, actor: str) -> ReceivableUploadSummary:
    summary = ReceivableUploadSummary()
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as error:  # noqa: BLE001 - surface any parse failure
        summary.errors.append(f"Could not read the Excel file: {error}")
        return summary

    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        summary.errors.append("The Excel file is empty.")
        return summary

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    for index, raw in enumerate(rows[1:], start=2):
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        try:
            invoice_value = _cell(row, "invoice value", "invoice_value", "amount")
            if invoice_value in (None, ""):
                summary.skipped += 1
                continue
            create_payable(
                CreatePayableRequest(
                    partner_type=str(_cell(row, "partner type", "partner_type") or "other").strip().lower().replace(" ", "_"),
                    partner_name=str(_cell(row, "partner name", "partner_name", "partner") or "").strip(),
                    country=str(_cell(row, "country") or "").strip(),
                    invoice_number=str(_cell(row, "invoice number", "invoice_number", "invoice") or "").strip(),
                    invoice_date=str(_cell(row, "invoice date", "invoice_date") or date.today().isoformat())[:10],
                    due_date=str(_cell(row, "due date", "due_date") or date.today().isoformat())[:10],
                    payment_terms=(str(_cell(row, "payment terms", "payment_terms")) if _cell(row, "payment terms", "payment_terms") else None),
                    invoice_value=float(invoice_value),
                    paid_value=float(_cell(row, "paid value", "paid_value") or 0),
                    actor=actor,
                )
            )
            summary.created += 1
        except (ValueError, TypeError) as error:
            summary.errors.append(f"Row {index}: {error}")
            summary.skipped += 1
    return summary
