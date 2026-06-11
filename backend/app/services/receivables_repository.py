"""Receivables Engine (Phase 3A, P1) + credit-limit store (used by P7).

Tracks distributor invoices, outstanding/paid values, payment history, and an
Excel bulk upload. Status (open / partially_paid / paid / overdue) is computed
from paid value and due date at read time. Every write is audited.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.db.local_persistence import load_collection, record_audit_event, save_collection
from app.schemas.receivables import (
    CreateReceivableRequest,
    CreditLimit,
    PaymentEntry,
    Receivable,
    ReceivableUploadSummary,
    RecordPaymentRequest,
    SetCreditLimitRequest,
)


def _load() -> list[Receivable]:
    return load_collection("receivables", lambda payload: Receivable(**payload))


def _save(records: list[Receivable]) -> None:
    save_collection("receivables", records, lambda record: record.receivable_id)


def _next_id(existing: list[Receivable]) -> str:
    numbers = [
        int(r.receivable_id.split("-")[-1])
        for r in existing
        if r.receivable_id.startswith("RCV-") and r.receivable_id.split("-")[-1].isdigit()
    ]
    return f"RCV-{(max(numbers) + 1) if numbers else 1:04d}"


def _effective_status(record: Receivable, today: date | None = None) -> str:
    today = today or date.today()
    outstanding = record.invoice_value - record.paid_value
    if outstanding <= 0:
        return "paid"
    try:
        overdue = date.fromisoformat(record.due_date) < today
    except ValueError:
        overdue = False
    if overdue:
        return "overdue"
    return "partially_paid" if record.paid_value > 0 else "open"


def _enrich(record: Receivable) -> Receivable:
    record.outstanding_value = max(record.invoice_value - record.paid_value, 0)
    record.status = _effective_status(record)
    return record


def create_receivable(request: CreateReceivableRequest) -> Receivable:
    if request.invoice_value <= 0:
        raise ValueError("Invoice value must be greater than zero.")
    records = _load()
    record = Receivable(
        receivable_id=_next_id(records),
        distributor=request.distributor.strip(),
        country=request.country.strip(),
        invoice_number=request.invoice_number.strip(),
        invoice_date=request.invoice_date,
        due_date=request.due_date,
        payment_terms=request.payment_terms,
        invoice_value=request.invoice_value,
        paid_value=request.paid_value,
        created_by=request.actor,
        updated_at=datetime.now().isoformat(),
    )
    _save([record, *records])
    record_audit_event(
        action="receivable_create",
        module_name="receivables",
        entity_name="receivable",
        entity_id=record.receivable_id,
        actor=request.actor,
        new_value=record,
    )
    return _enrich(record)


def list_receivables(distributor: str | None = None, status: str | None = None, country: str | None = None) -> list[Receivable]:
    records = [_enrich(r) for r in _load()]
    if distributor:
        records = [r for r in records if r.distributor.lower() == distributor.lower()]
    if country:
        records = [r for r in records if r.country.lower() == country.lower()]
    if status:
        records = [r for r in records if r.status == status.lower()]
    return sorted(records, key=lambda r: r.invoice_date, reverse=True)


def get_receivable(receivable_id: str) -> Receivable | None:
    record = next((r for r in _load() if r.receivable_id == receivable_id), None)
    return _enrich(record) if record else None


def record_payment(receivable_id: str, request: RecordPaymentRequest) -> Receivable:
    if request.amount <= 0:
        raise ValueError("Payment amount must be greater than zero.")
    records = _load()
    record = next((r for r in records if r.receivable_id == receivable_id), None)
    if not record:
        raise ValueError(f"Receivable not found: {receivable_id}")
    record.paid_value = min(record.invoice_value, record.paid_value + request.amount)
    record.payment_history.append(
        PaymentEntry(amount=request.amount, paid_date=request.paid_date or date.today().isoformat(), note=request.note, actor=request.actor)
    )
    record.updated_at = datetime.now().isoformat()
    _save(records)
    record_audit_event(
        action="receivable_payment",
        module_name="receivables",
        entity_name="receivable",
        entity_id=receivable_id,
        actor=request.actor,
        new_value=record,
    )
    return _enrich(record)


def _cell(row: dict, *names):
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def upload_receivables(content: bytes, actor: str) -> ReceivableUploadSummary:
    summary = ReceivableUploadSummary()
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as error:  # noqa: BLE001 - report any parse failure to the caller
        summary.errors.append(f"Could not read the Excel file: {error}")
        return summary

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        summary.errors.append("The Excel file is empty.")
        return summary

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    for index, raw in enumerate(rows[1:], start=2):
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        try:
            invoice_value = _cell(row, "invoice value", "invoice_value", "amount")
            if invoice_value in (None, ""):
                summary.skipped += 1
                continue
            create_receivable(
                CreateReceivableRequest(
                    distributor=str(_cell(row, "distributor") or "").strip(),
                    country=str(_cell(row, "country") or "").strip(),
                    invoice_number=str(_cell(row, "invoice number", "invoice_number", "invoice") or "").strip(),
                    invoice_date=str(_cell(row, "invoice date", "invoice_date") or date.today().isoformat())[:10],
                    due_date=str(_cell(row, "due date", "due_date") or date.today().isoformat())[:10],
                    payment_terms=(str(_cell(row, "payment terms", "payment_terms")) if _cell(row, "payment terms", "payment_terms") else None),
                    invoice_value=float(invoice_value),
                    paid_value=float(_cell(row, "paid value", "paid_value") or 0),
                    actor=actor,
                )
            )
            summary.created += 1
        except (ValueError, TypeError) as error:
            summary.errors.append(f"Row {index}: {error}")
            summary.skipped += 1
    return summary


# --- credit-limit store (used by P7 credit control) ---------------------

def _load_credit_limits() -> list[CreditLimit]:
    return load_collection("credit_limits", lambda payload: CreditLimit(**payload))


def _save_credit_limits(records: list[CreditLimit]) -> None:
    save_collection("credit_limits", records, lambda record: record.distributor.lower())


def get_credit_limit(distributor: str) -> CreditLimit | None:
    return next((c for c in _load_credit_limits() if c.distributor.lower() == distributor.lower()), None)


def set_credit_limit(request: SetCreditLimitRequest) -> CreditLimit:
    records = [c for c in _load_credit_limits() if c.distributor.lower() != request.distributor.lower()]
    existing = get_credit_limit(request.distributor)
    record = CreditLimit(
        distributor=request.distributor.strip(),
        credit_limit=request.credit_limit,
        override=existing.override if existing else False,
        override_reason=existing.override_reason if existing else None,
        updated_by=request.actor,
        updated_at=datetime.now().isoformat(),
    )
    _save_credit_limits([record, *records])
    record_audit_event(
        action="credit_limit_set",
        module_name="credit_control",
        entity_name="credit_limit",
        entity_id=record.distributor,
        actor=request.actor,
        new_value=record,
    )
    return record


def set_credit_override(distributor: str, override: bool, reason: str | None, actor: str) -> CreditLimit:
    existing = get_credit_limit(distributor)
    records = [c for c in _load_credit_limits() if c.distributor.lower() != distributor.lower()]
    record = CreditLimit(
        distributor=distributor.strip(),
        credit_limit=existing.credit_limit if existing else 0,
        override=override,
        override_reason=reason,
        updated_by=actor,
        updated_at=datetime.now().isoformat(),
    )
    _save_credit_limits([record, *records])
    record_audit_event(
        action="credit_override",
        module_name="credit_control",
        entity_name="credit_limit",
        entity_id=record.distributor,
        actor=actor,
        new_value=record,
    )
    return record
